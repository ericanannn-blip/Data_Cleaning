"""Content readers, imported only in a disposable, resource-bounded process.

Uploaded text is data. No eval, SQL supplied by users, macros, or document
instructions are executed. Extension and declared MIME are never trusted.
"""

from __future__ import annotations

import bz2
import codecs
import csv
import gzip
import io
import itertools
import json
import lzma
import math
import re
import shutil
import sqlite3
import subprocess
import sys
import tempfile
import zipfile
from datetime import date, datetime
from pathlib import Path

from .inspection import reader
from .schemas import infer_schemas

READER_VERSION = "3.0"


def safe_value(value, depth=0):
    if depth > 5:
        return "[depth limited]"
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, float) and not math.isfinite(value):
        return None
    if isinstance(value, bytes):
        return {"bytes": len(value), "prefix_hex": value[:32].hex()}
    if isinstance(value, dict):
        return {str(k)[:256]: safe_value(v, depth + 1) for k, v in itertools.islice(value.items(), 512)}
    if isinstance(value, (list, tuple)):
        return [safe_value(v, depth + 1) for v in value[:200]]
    if isinstance(value, str):
        return value[:2048]
    if value is None or isinstance(value, (bool, int, float)):
        return value
    return str(value)[:2048]


def decode_sample(data, complete=False):
    encodings = ("utf-32",) if data.startswith((b"\x00\x00\xfe\xff", b"\xff\xfe\x00\x00")) else (
        "utf-16",) if data.startswith((b"\xff\xfe", b"\xfe\xff")) else (
        "utf-8-sig",) if data.startswith(b"\xef\xbb\xbf") else ("utf-8", "gb18030")
    for encoding in encodings:
        try:
            # A sampled window can end within a character. Buffer that partial
            # character; never strip arbitrary binary bytes to manufacture text.
            text = codecs.getincrementaldecoder(encoding)().decode(data, final=complete)
            if text and reader.is_readable(text):
                return text, encoding
        except UnicodeError:
            continue
    return None


def spread(items, budget):
    if len(items) <= budget:
        return items
    if budget == 1:
        return [items[len(items) // 2]]
    return [items[round(index * (len(items) - 1) / (budget - 1))] for index in range(budget)]


def windows(path, budget):
    size = path.stat().st_size
    if size <= budget:
        with path.open("rb") as stream:
            return [(0, stream.read(budget))]
    widths = [budget // 2, budget // 4, budget - budget // 2 - budget // 4]
    offsets = [0, max(0, size // 2 - widths[1] // 2), max(0, size - widths[2])]
    result = []
    with path.open("rb") as stream:
        for offset, width in zip(offsets, widths):
            # Align UTF-16/32 code units too; no significance for other encodings.
            offset -= offset % 4
            stream.seek(offset)
            result.append((offset, stream.read(width)))
    return result


def text_records(path, options, warnings):
    parts = windows(path, options["bytes"])
    decoded = decode_sample(parts[0][1], complete=path.stat().st_size <= options["bytes"])
    if not decoded:
        return None
    first, encoding = decoded
    window_encoding = encoding
    if encoding == "utf-16":
        window_encoding = "utf-16-be" if parts[0][1].startswith(b"\xfe\xff") else "utf-16-le"
    elif encoding == "utf-32":
        window_encoding = "utf-32-be" if parts[0][1].startswith(b"\x00\x00\xfe\xff") else "utf-32-le"
    texts = [first]
    for _, data in parts[1:]:
        text = data.decode(window_encoding, errors="replace")
        lines = text.splitlines()
        texts.append("\n".join(lines[1:-1]))
    limit = options["rows"]
    stripped = first.lstrip()
    if stripped.startswith("{\\rtf"):
        from striprtf.striprtf import rtf_to_text
        return "rtf", "document", [{"text": rtf_to_text(first)[:2048]}], {"encoding": encoding}, parts
    if stripped.startswith(("{", "[")):
        import ijson
        records = []
        # JSON lines are sampled at several byte positions, unlike JSON arrays.
        lines = first.splitlines()
        if len(lines) > 1 and all(line.lstrip().startswith("{") for line in lines[:3] if line.strip()):
            failed = 0
            for text in texts:
                for line in spread(text.splitlines(), max(1, limit // len(texts))):
                    try:
                        value = json.loads(line)
                        records.append(safe_value(value) if isinstance(value, dict) else {"value": safe_value(value)})
                    except ValueError:
                        failed += 1
            if failed:
                warnings.append(f"{failed} sampled JSONL lines could not be parsed (including possible cut boundaries)")
            return "jsonl", "tabular", records[:limit], {"encoding": encoding}, parts
        with path.open("r", encoding=encoding) as text_stream:
            # ijson consumes bytes; transcode non-UTF8 inputs incrementally.
            class EncodedStream:
                def read(self, amount=-1):
                    return text_stream.read(amount).encode("utf-8")
            stream = EncodedStream()
            try:
                if stripped.startswith("["):
                    values = itertools.islice(ijson.items(stream, "item", use_float=True), limit)
                    records = [safe_value(v) if isinstance(v, dict) else {"value": safe_value(v)} for v in values]
                elif path.stat().st_size <= options["bytes"]:
                    value = json.load(text_stream)
                    array = next((value[k] for k in ("records", "rows", "data", "events", "items")
                                  if isinstance(value.get(k), list)), None)
                    records = ([safe_value(v) if isinstance(v, dict) else {"value": safe_value(v)} for v in array[:limit]]
                               if array is not None else [safe_value(value)])
                else:
                    match = re.search(r'"(records|rows|data|events|items)"\s*:\s*\[', first)
                    if match:
                        values = itertools.islice(ijson.items(stream, match[1] + ".item", use_float=True), limit)
                        records = [safe_value(v) if isinstance(v, dict) else {"value": safe_value(v)} for v in values]
                    else:
                        pairs = itertools.islice(ijson.kvitems(stream, "", use_float=True), min(limit, 512))
                        records = [{str(k): safe_value(v) for k, v in pairs}]
                warnings.append("JSON records sampled from the prefix; the unsampled suffix was not validated")
                return "json", "tabular", records, {"encoding": encoding}, parts[:1]
            except (ValueError, ijson.JSONError) as exc:
                warnings.append("JSON-like content has invalid sampled syntax: " + type(exc).__name__)
                return "json", "document", [{"text": first[:2048]}], {"encoding": encoding, "invalid_syntax": True}, parts
    if stripped.startswith("<"):
        if re.match(r"(?is)^\s*(<!doctype\s+html|<html\b)", first):
            from html.parser import HTMLParser
            class TextParser(HTMLParser):
                def __init__(self):
                    super().__init__()
                    self.pieces, self.hidden = [], 0
                def handle_starttag(self, tag, attrs):
                    if tag in {"script", "style"}:
                        self.hidden += 1
                def handle_endtag(self, tag):
                    if tag in {"script", "style"}:
                        self.hidden = max(0, self.hidden - 1)
                def handle_data(self, data):
                    if not self.hidden:
                        self.pieces.append(data)
            parser = TextParser()
            parser.feed(first)
            return "html", "document", [{"text": " ".join(parser.pieces)[:2048]}], {"encoding": encoding}, parts
        if path.stat().st_size <= options["bytes"]:
            from defusedxml import ElementTree
            root = ElementTree.fromstring(first, forbid_dtd=True)
            records = []
            for child in itertools.islice(root, limit):
                record = {"tag": child.tag, **child.attrib}
                record.update({node.tag: (node.text or "").strip() for node in itertools.islice(child, 512)})
                if not len(child):
                    record["text"] = child.text or ""
                records.append(safe_value(record))
            return "xml", "tabular", records or [{"text": " ".join(root.itertext())[:2048]}], {"encoding": encoding}, parts
        warnings.append("Large XML was sampled as text; structure needs a streaming domain-specific adapter")
        return "xml", "document", [{"text": t[:2048]} for t in texts], {"encoding": encoding}, parts
    try:
        dialect = csv.Sniffer().sniff(first[:8192], delimiters=",\t;|")
    except csv.Error:
        dialect = None
        lines = first.splitlines()[:20]
        if lines:
            choices = [delimiter for delimiter in ",\t;|" if lines[0].count(delimiter) and
                       sum(delimiter in line for line in lines) >= max(2, len(lines) // 2)]
            if choices:
                delimiter = max(choices, key=lines[0].count)
                dialect = type("SampleDialect", (csv.excel,), {"delimiter": delimiter})
                warnings.append("Delimiter inferred despite inconsistent row widths")
    numeric = re.compile(r"^[+-]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][+-]?\d+)?$")
    whitespace = not dialect and any(len(cells := line.split()) > 1 and all(numeric.fullmatch(cell) for cell in cells)
                                    for line in first.splitlines()[:20])
    if dialect or whitespace:
        def rows_for(text):
            if whitespace:
                return [re.split(r"\s+", line.strip()) for line in text.splitlines() if line.strip()]
            return list(csv.reader(io.StringIO(text), dialect))
        rows = rows_for(first)
        if len(rows) > 1:
            is_header = any(not numeric.fullmatch(cell.strip()) for cell in rows[0])
            header = rows.pop(0) if is_header else [f"column_{i + 1}" for i in range(len(rows[0]))]
            unique = []
            for i, name in enumerate(header):
                base = name.strip() or f"column_{i + 1}"
                candidate, suffix = base, 2
                while candidate in unique:
                    candidate, suffix = f"{base}_{suffix}", suffix + 1
                unique.append(candidate)
            if unique != header:
                warnings.append("Blank or duplicate source headers were disambiguated; original headers are retained in metadata")
            records, ragged = [], 0
            for index, text in enumerate(texts):
                selected = rows if index == 0 else rows_for(text)
                for row in spread(selected, max(1, limit // len(texts))):
                    if len(row) != len(unique):
                        ragged += 1
                    record = dict(zip(unique, row))
                    if len(row) > len(unique):
                        record["_extra_cells"] = row[len(unique):]
                    records.append(safe_value(record))
            if ragged:
                warnings.append(f"{ragged} sampled rows have a different width from the header")
            return "table", "tabular", records[:limit], {"encoding": encoding, "source_headers": header,
                    "delimiter": "whitespace" if whitespace else dialect.delimiter, "header_inferred": is_header}, parts
    return "text", "document", [{"text": text[:2048], "sample_offset": offset}
            for (offset, _), text in zip(parts, texts)], {"encoding": encoding}, parts


def ocr_image(image, options, warnings):
    if not shutil.which("tesseract"):
        warnings.append("Tesseract is unavailable; image text needs OCR")
        return "", "needs_ocr", {}
    with tempfile.TemporaryDirectory() as temporary:
        target = Path(temporary) / "sample.png"
        (image if image.mode in {"1", "L", "RGB", "RGBA"} else image.convert("RGB")).save(target)
        process = subprocess.run(["tesseract", str(target), "stdout", "-l", options["ocr_languages"], "tsv"],
                                 capture_output=True, timeout=max(1, options.get("timeout", 60) - 5), check=False)
    if process.returncode:
        warnings.append("Tesseract could not extract text; check image integrity and configured language models")
        return "", "needs_ocr", {}
    words, confidences = [], []
    for row in csv.DictReader(io.StringIO(process.stdout.decode("utf-8", "replace")), delimiter="\t"):
        text = (row.get("text") or "").strip()
        try:
            confidence = float(row.get("conf", -1))
        except ValueError:
            continue
        if text and confidence >= 0:
            words.append(text)
            confidences.append(confidence)
        if len(words) >= 2000:
            warnings.append("OCR word sample capped at 2,000 words")
            break
    mean = round(sum(confidences) / len(confidences), 2) if confidences else None
    diagnostics = {"ocr_words_sampled": len(words), "ocr_mean_word_confidence": mean,
                   "ocr_confidence_scope": "engine score on sampled words, not a calibrated accuracy estimate"}
    if mean is not None and mean < 60:
        warnings.append("Low OCR engine confidence; content requires review or a stronger reader")
        status = "needs_review"
    else:
        status = "sampled"
    if not words:
        warnings.append("No recognizable text found; image semantics were not interpreted")
    return " ".join(words)[:20000], status, diagnostics


def office_records(path, fmt, options, warnings):
    from defusedxml import ElementTree
    records = []
    with zipfile.ZipFile(path) as archive:
        if fmt == "xlsx":
            import openpyxl
            with path.open("rb") as source:
                workbook = openpyxl.load_workbook(source, read_only=True, data_only=False, keep_links=False)
                try:
                    for sheet in workbook.worksheets[:options["pages"]]:
                        iterator = sheet.iter_rows(values_only=True)
                        header = [str(v) if v is not None else f"column_{i + 1}"
                                  for i, v in enumerate(next(iterator, ()))]
                        # Duplicate columns remain distinct, just as for CSV.
                        counts = {}
                        for i, name in enumerate(header):
                            counts[name] = counts.get(name, 0) + 1
                            if counts[name] > 1:
                                header[i] = f"{name}_{counts[name]}"
                                warnings.append("Duplicate workbook headers were disambiguated")
                        for row in itertools.islice(iterator, max(1, options["rows"] // options["pages"])):
                            records.append(safe_value({"_sheet": sheet.title, **dict(zip(header, row))}))
                finally:
                    workbook.close()
            warnings.append("Workbook formulas are preserved as strings and are never executed")
            return records
        names = (["word/document.xml"] if fmt == "docx" else
                 sorted(n for n in archive.namelist() if re.fullmatch(r"ppt/slides/slide\d+\.xml", n))[:options["pages"]]
                 if fmt == "pptx" else ["content.xml"])
        consumed = 0
        for name in names:
            with archive.open(name) as stream:
                for _, element in ElementTree.iterparse(stream, events=("end",), forbid_dtd=True):
                    text = (element.text or "").strip()
                    if text:
                        consumed += len(text.encode("utf-8"))
                        records.append({"section": name, "text": text[:2048]})
                    element.clear()
                    if consumed >= options["bytes"] or len(records) >= options["rows"]:
                        break
            if consumed >= options["bytes"] or len(records) >= options["rows"]:
                break
    return records


def extract(path: Path, options: dict, depth=0) -> dict:
    warnings, records, metadata = [], [], {}
    size = path.stat().st_size
    with path.open("rb") as stream:
        prefix = stream.read(min(options["bytes"], 65536))
    signature = reader.detect_format(prefix)
    fmt, family, status = signature[0] if signature else "unknown", "opaque", "sampled"
    # Pillow's plugin registry covers many more image encodings than the small
    # legacy magic table, and does not trust the .dat extension.
    if prefix and not signature and not decode_sample(prefix, complete=size <= len(prefix)) and not prefix.startswith((b"PAR1", b"\x93NUMPY", b"\x89HDF\r\n\x1a\n")):
        from PIL import Image, UnidentifiedImageError
        try:
            with Image.open(path) as image:
                fmt = "pillow_image"
        except (UnidentifiedImageError, OSError, ValueError):
            pass
    if not signature and prefix.startswith(b"\x1aE\xdf\xa3"):
        fmt = "matroska"
    elif not signature and prefix.startswith(b"FLV"):
        fmt = "flv"
    elif not signature and prefix[:4] == b"RIFF" and prefix[8:12] == b"AVI ":
        fmt = "avi"
    sampling = {"method": "bounded prefix", "bytes_budget": options["bytes"],
                "rows_budget": options["rows"], "pages_budget": options["pages"],
                "file_size": size, "full_file_validation": False}
    if not prefix:
        status = "empty"
    elif fmt == "pdf":
        from pypdf import PdfReader
        with path.open("rb") as stream:
            document = PdfReader(stream, strict=False)
            if document.is_encrypted:
                return result("pdf", "document", "encrypted", {}, [], sampling, ["A password is required"], options)
            total = len(document.pages)
            pages = sorted({round(i * (total - 1) / max(1, min(options["pages"], total) - 1))
                            for i in range(min(options["pages"], total))})
            metadata["page_count"] = total
            sampling.update(method="first/middle/last pages", page_indices=pages)
            for index in pages:
                text = document.pages[index].extract_text() or ""
                if not text.strip():
                    try:
                        import pypdfium2
                        with pypdfium2.PdfDocument(path) as rendered:
                            page = rendered[index]
                            width, height = page.get_size()
                            scale = min(2.0, (20_000_000 / max(1, width * height)) ** .5)
                            bitmap = page.render(scale=scale)
                            try:
                                text, ocr_status, diagnostics = ocr_image(bitmap.to_pil(), options, warnings)
                                metadata.setdefault("ocr_pages", []).append({"page": index + 1, **diagnostics})
                                if ocr_status != "sampled":
                                    status = ocr_status
                            finally:
                                bitmap.close()
                                page.close()
                    except ImportError:
                        warnings.append("Scanned PDF requires pypdfium2 and Tesseract")
                        status = "needs_ocr"
                records.append({"page": index + 1, "text": text[:20000]})
            family = "document"
    elif fmt in {"png", "jpeg", "gif", "webp", "bmp", "pillow_image"}:
        from PIL import Image
        with Image.open(path) as checked:
            checked.verify()
        with Image.open(path) as image:
            fmt, family = image.format.lower(), "media_asset"
            metadata = {"width": image.width, "height": image.height, "mode": image.mode,
                        "frames": getattr(image, "n_frames", 1)}
            text, status, diagnostics = ocr_image(image, options, warnings)
            metadata.update(diagnostics)
            records = [{**metadata, "ocr_text": text}]
            sampling.update(method="first image frame and OCR", frames_sampled=1)
    elif fmt == "zip":
        with zipfile.ZipFile(path) as archive:
            names = set(archive.namelist())
            fmt = ("docx" if "word/document.xml" in names else "xlsx" if "xl/workbook.xml" in names else
                   "pptx" if "ppt/presentation.xml" in names else "odf" if "content.xml" in names else "zip")
            if fmt == "zip":
                records = [{"member": info.filename, "size_bytes": info.file_size,
                            "compressed_bytes": info.compress_size, "encrypted": bool(info.flag_bits & 1)}
                           for info in itertools.islice(archive.infolist(), options["rows"])]
                metadata["member_count"] = len(archive.infolist())
                family = "archive_inventory"
                warnings.append("Archive inventory sampled; members require explicit import --expand for content reading")
        if fmt != "zip":
            family = "tabular" if fmt == "xlsx" else "document"
            records = office_records(path, fmt, options, warnings)
    elif fmt in {"gzip", "bzip2", "xz", "zlib"}:
        if depth >= 3:
            return result(fmt, "opaque", "resource_limit", {}, [], sampling, ["Compression nesting limit reached"], options)
        opener = {"gzip": gzip.open, "bzip2": bz2.open, "xz": lzma.open}.get(fmt)
        with tempfile.TemporaryDirectory() as temporary:
            decoded = Path(temporary) / "decoded.dat"
            if opener:
                with opener(path, "rb") as source:
                    data = source.read(options["bytes"] + 1)
            else:
                import zlib
                data = zlib.decompressobj().decompress(prefix, options["bytes"] + 1)
            decoded.write_bytes(data[:options["bytes"]])
            inner = extract(decoded, options, depth + 1)
            inner["container_format"] = fmt
            inner["sampling"]["file_size"] = size
            inner["sampling"]["method"] = "bounded decompressed prefix"
            inner["warnings"].append("Compressed stream sampled within byte budget; unsampled integrity was not checked")
            return inner
    elif fmt == "sqlite":
        family = "tabular"
        with sqlite3.connect(path.as_uri() + "?mode=ro&immutable=1", uri=True) as connection:
            connection.execute("PRAGMA query_only=ON")
            connection.execute("PRAGMA trusted_schema=OFF")
            tables = connection.execute("SELECT name FROM sqlite_master WHERE type='table' AND name NOT LIKE 'sqlite_%' LIMIT ?",
                                        (options["pages"],)).fetchall()
            metadata["sampled_tables"] = [row[0] for row in tables]
            for (name,) in tables:
                quoted = '"' + name.replace('"', '""') + '"'
                cursor = connection.execute(f"SELECT * FROM {quoted} LIMIT ?", (max(1, options["rows"] // max(1, len(tables))),))
                header = [col[0] for col in cursor.description]
                records.extend(safe_value({"_table": name, **dict(zip(header, row))}) for row in cursor)
            warnings.append("SQLite tables sampled read-only; triggers and supplied queries are not executed")
    elif fmt in {"mp3", "mp4", "wav", "ogg", "flac", "matroska", "flv", "avi"}:
        family = "media_asset"
        if shutil.which("ffprobe"):
            process = subprocess.run(["ffprobe", "-v", "error", "-show_format", "-show_streams", "-of", "json", str(path)],
                                     capture_output=True, timeout=max(1, options.get("timeout", 60) - 5), check=False)
            if process.returncode:
                status = "corrupt"
            else:
                metadata = safe_value(json.loads(process.stdout))
                metadata.get("format", {}).pop("filename", None)
                records = metadata.get("streams", [])
                warnings.append("Media structure and technical metadata read; speech and video semantics require a transcription/vision adapter")
        else:
            status = "dependency_missing"
            warnings.append("ffprobe is required to inspect media")
    elif prefix.startswith(b"PAR1"):
        fmt, family = "parquet", "tabular"
        import pyarrow.parquet as parquet
        source = parquet.ParquetFile(path)
        metadata = {"total_rows": source.metadata.num_rows, "row_groups": source.num_row_groups}
        records = [safe_value(row) for batch in itertools.islice(source.iter_batches(batch_size=options["rows"]), 1)
                   for row in batch.to_pylist()]
    elif prefix.startswith(b"\x93NUMPY"):
        fmt, family = "npy", "array"
        import numpy as np
        array = np.load(path, mmap_mode="r", allow_pickle=False)
        metadata = {"shape": list(array.shape), "dtype": str(array.dtype)}
        records = [{"index": i, "value": safe_value(array.flat[i].item())}
                   for i in spread(range(array.size), options["rows"])]
        sampling.update(method="systematic positions across the array")
    elif prefix.startswith(b"\x89HDF\r\n\x1a\n"):
        fmt, family = "hdf5", "array"
        import h5py
        with h5py.File(path, "r") as source:
            def visit(name, obj):
                if isinstance(obj, h5py.Dataset) and len(records) < options["rows"]:
                    structure = {"dataset": name, "shape": list(obj.shape), "dtype": str(obj.dtype)}
                    if obj.size == 0:
                        records.append(structure)
                        return
                    count = min(16, obj.size, options["rows"] - len(records))
                    for flat in spread(range(obj.size), count):
                        import numpy as np
                        index = tuple(int(i) for i in np.unravel_index(flat, obj.shape)) if obj.shape else ()
                        value = obj[index]
                        if hasattr(value, "tolist"):
                            value = value.tolist()
                        records.append(safe_value({**structure, "index": list(index), "value": value}))
            source.visititems(visit)
        sampling.update(method="systematic positions across sampled datasets")
        warnings.append("HDF5 datasets and bounded scalar values sampled; unsampled datasets and values were not validated")
    else:
        text = text_records(path, options, warnings)
        if text:
            fmt, family, records, metadata, parts = text
            sampling.update(method="first/middle/last byte windows" if len(parts) > 1 else "bounded prefix",
                            byte_offsets=[offset for offset, _ in parts], bytes_read=sum(len(data) for _, data in parts))
            if metadata.get("invalid_syntax"):
                status = "corrupt"
        elif fmt in {"7z", "rar"}:
            status = "dependency_missing"
            warnings.append("This archive requires an optional format-specific reader")
        else:
            # Restore the existing reader's single-byte XOR images in bounded memory.
            with path.open("rb") as stream:
                stream.seek(max(0, size - 2))
                ending = stream.read(2)
            xor = reader.find_image_xor(prefix if size <= len(prefix) else prefix + ending)
            if xor:
                with tempfile.TemporaryDirectory() as temporary:
                    target = Path(temporary) / "decoded.dat"
                    with path.open("rb") as source, target.open("wb") as output:
                        while chunk := source.read(65536):
                            output.write(bytes(byte ^ xor[1] for byte in chunk))
                    inner = extract(target, options, depth + 1)
                    inner["transform"] = [f"xor:{xor[1]}"]
                    return inner
            status = "unsupported"
            metadata["header_hex"] = prefix[:32].hex()
            warnings.append("Unknown or proprietary binary content; original preserved for a new adapter")
    if status in {"unsupported", "dependency_missing"} and options.get("tika_url"):
        import urllib.error
        import urllib.request
        from urllib.parse import urlsplit
        url = urlsplit(options["tika_url"])
        if url.scheme not in {"http", "https"} or not url.hostname:
            raise ValueError("Invalid TIKA_URL")
        try:
            with path.open("rb") as source:
                request = urllib.request.Request(options["tika_url"].rstrip("/") + "/tika", method="PUT",
                    data=iter(lambda: source.read(65536), b""), headers={"Content-Type": "application/octet-stream",
                        "Accept": "text/plain", "Content-Length": str(size)})
                # Standard urllib honors inherited proxy and CA configuration.
                with urllib.request.urlopen(request, timeout=max(1, options["timeout"] - 5)) as response:
                    text = response.read(options["bytes"]).decode("utf-8", "replace")
                    if text.strip():
                        records, family, status = [{"text": text[:2048]}], "document", "sampled"
                        metadata["fallback_reader"] = "apache_tika"
                        warnings.append("Operator-configured Tika fallback used; response sampled, server has separate resource requirements")
        except OSError:
            warnings.append("Configured Tika service is unavailable; original reader status retained")
    if family == "document" and not any(str(r.get("text", "")).strip() for r in records) and status == "sampled":
        status = "needs_ocr" if fmt == "pdf" else "empty_content"
    return result(fmt, family, status, metadata, records, sampling, warnings, options)


def result(fmt, family, status, metadata, records, sampling, warnings, options):
    records = [safe_value(record) for record in records[:options["rows"]]]
    sampling["records_sampled"] = len(records)
    schema = infer_schemas(records, family, sampling, warnings) if records else {"candidates": [], "sampling": sampling}
    return {"reader_version": READER_VERSION, "format": fmt, "family": family,
            "status": status, "metadata": safe_value(metadata), "sample": records[:5],
            "sampling": sampling, "warnings": warnings, "schema": schema}


def main():
    options = json.loads(sys.argv[2])
    # Set limits before importing any native parser. Windows uses the timeout
    # only; deploy the Linux container for enforced memory/CPU limits.
    try:
        import resource
        memory = options["memory_mb"] * 1024 * 1024
        resource.setrlimit(resource.RLIMIT_AS, (memory, memory))
        resource.setrlimit(resource.RLIMIT_CPU, (options["timeout"], options["timeout"] + 1))
    except ImportError:
        pass
    try:
        analysis = extract(Path(sys.argv[1]).resolve(), options)
    except ImportError as exc:
        analysis = {"reader_version": READER_VERSION, "status": "dependency_missing",
                    "error_code": exc.name, "schema": {"candidates": []}}
    except MemoryError:
        analysis = {"reader_version": READER_VERSION, "status": "resource_limit", "error_code": "memory_limit", "schema": {"candidates": []}}
    except Exception as exc:
        status = "resource_limit" if isinstance(exc, subprocess.TimeoutExpired) or type(exc).__name__ == "DecompressionBombError" else "corrupt"
        analysis = {"reader_version": READER_VERSION, "status": status, "error_code": type(exc).__name__, "schema": {"candidates": []}}
    print(json.dumps(analysis, ensure_ascii=False, allow_nan=False))


if __name__ == "__main__":
    main()
